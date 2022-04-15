

#from docx import Document

from mailmerge import MailMerge

import openpyxl, docx



#Load Workbook

wk = openpyxl.load_workbook(r"C:\Users\shill\PycharmProjects\Excel\TestWritePyExcel.xlsx")


sheets = wk.sheetnames

print(sheets[0])


sh = wk[sheets[0]]

Speaker1 = sh['B3'].value

doc = docx.Document('WordPython.docx')

print(doc.paragraphs)

doc.add_paragraph(Speaker1)




doc.save('WordPython.docx')

paro = doc.paragraphs

for par in paro:
    print(par)